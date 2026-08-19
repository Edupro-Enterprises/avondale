import openpyxl
wb = openpyxl.load_workbook("/home/frappe/frappe-bench/Grade7_Pastel_Balances_Clean.xlsx", data_only=True)
for sheet_name in wb.sheetnames:
    ws = wb[sheet_name]
    print(f"=== SHEET: {sheet_name} (dims: {ws.dimensions}, max_row={ws.max_row}, max_col={ws.max_column}) ===")
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=min(15, ws.max_row), values_only=True)):
        print(i, row)
